"""Table loading and plot preparation, independent of the desktop interface."""

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
import csv
import io
import math


class TableError(ValueError):
    pass


@dataclass
class Table:
    headers: list[str]
    rows: list[tuple]

    @property
    def numeric_columns(self):
        return [i for i in range(len(self.headers))
                if any(as_number(row[i]) is not None for row in self.rows)]


def as_number(value):
    if value is None or isinstance(value, (bool, date, datetime)):
        return None
    try:
        number = float(str(value).strip())
        return number if math.isfinite(number) else None
    except (ValueError, TypeError):
        return None


def make_table(values):
    iterator = iter(values)
    header = next(iterator, None)
    if header is None:
        raise TableError("表格为空，请选择包含列名和数据的表格。")
    rows = [tuple(row) for row in iterator if any(v is not None and str(v).strip() for v in row)]
    width = max([len(header)] + [len(row) for row in rows])
    while width and all(len(row) <= width - 1 or row[width - 1] in (None, "")
                        for row in [header, *rows]):
        width -= 1
    if not width or not rows:
        raise TableError("表格没有可绘制的数据，第一行应为列名，第二行起为数据。")
    headers, used = [], set()
    for i in range(width):
        original = str(header[i]).strip() if i < len(header) and header[i] is not None else ""
        base = original or f"列 {i + 1}"
        name, suffix = base, 2
        while name in used:
            name = f"{base} ({suffix})"
            suffix += 1
        used.add(name)
        headers.append(name)
    return Table(headers, [tuple(row[i] if i < len(row) else None for i in range(width)) for row in rows])


def sheet_names(path):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            return book.sheetnames
        finally:
            book.close()
    if path.suffix.lower() == ".xls":
        import xlrd
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            return book.sheet_names()
        finally:
            book.release_resources()
    return []


def load_table(path, sheet=None):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            return make_table(book[sheet or book.sheetnames[0]].values)
        finally:
            book.close()
    if suffix == ".xls":
        import xlrd
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            page = book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)
            values = []
            for r in range(page.nrows):
                row = []
                for cell in page.row(r):
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                    elif cell.ctype in (xlrd.XL_CELL_ERROR, xlrd.XL_CELL_BOOLEAN):
                        row.append(None)
                    else:
                        row.append(cell.value)
                values.append(row)
            return make_table(values)
        finally:
            book.release_resources()
    if suffix not in (".csv", ".tsv", ".txt"):
        raise TableError("不支持此文件格式，请选择 XLSX、XLS、CSV 或 TSV 文件。")
    raw = path.read_bytes()
    content = None
    encodings = ("utf-16",) if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "gb18030")
    for encoding in encodings:
        try:
            content = raw.decode(encoding)
            break
        except UnicodeError:
            pass
    if content is None:
        raise TableError("无法识别文件编码，请将文件另存为 UTF-8 CSV。")
    try:
        dialect = csv.Sniffer().sniff(content[:16384], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
    return make_table(csv.reader(io.StringIO(content), dialect))


def prepare_series(table, x_column, y_columns):
    """Preserve row order and represent invalid Y cells as gaps."""
    if not y_columns:
        raise TableError("请至少勾选一条折线。")
    if x_column is None:
        indices = list(range(len(table.rows)))
        x_values = list(range(1, len(table.rows) + 1))
        kind = "number"
    else:
        available = [(i, row[x_column]) for i, row in enumerate(table.rows)
                     if row[x_column] is not None and str(row[x_column]).strip()]
        if not available:
            raise TableError("所选 X 轴列没有有效数据。")
        numbers = [as_number(value) for _, value in available]
        if any(value is not None for value in numbers):
            valid = [(i, value) for (i, _), value in zip(available, numbers) if value is not None]
            kind = "number"
        elif all(isinstance(value, (datetime, date)) for _, value in available):
            valid, kind = available, "date"
        else:
            valid, kind = [(i, str(value)) for i, value in available], "category"
        indices, x_values = map(list, zip(*valid))
    series = []
    for column in y_columns:
        numbers = [as_number(table.rows[i][column]) for i in indices]
        if not any(value is not None for value in numbers):
            raise TableError(f"“{table.headers[column]}”在有效 X 轴行中没有数值，请取消勾选该列。")
        series.append([value if value is not None else math.nan for value in numbers])
    dropped = len(table.rows) - len(indices)
    missing = sum(sum(math.isnan(value) for value in values) for values in series)
    return x_values, series, kind, dropped, missing
